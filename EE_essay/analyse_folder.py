#!/usr/bin/env python3
"""
analyse_folder.py -- batch analysis of Tracker exports for the coupled magnetic-mechanical oscillator EE.

    python3 analyse_folder.py "Tracker EE txt" --out results

Folder layout expected (anything else is skipped with a warning):
    Tracker EE txt/24.5/24.5_1.txt, 24.5_2.txt, ...      (d in mm = folder name; trial = number after '_')
    Tracker EE txt/26.5/...  Tracker EE txt/31/...  Tracker EE txt/Infinity/Infinity_1.txt (single spring / far apart)
Each txt is a Tracker export ("#multi:" header, then the track names, then 't x x' or 't x y x y', tab separated).
Blank cells (autotracker stopped early on one mass) and different end times are handled automatically.

For every trial:  release detection (flat hold segment removed) -> means removed -> normal coordinates
    eta = xA + xB (in-phase, contains only f1)  and  xi = xB - xA (anti-phase, contains only f2)
    -> f1, f2 from the single peaks of FFT(eta), FFT(xi)  (no mutual leakage; works even when the two peaks
       in xA alone are closer than one frequency bin)
    -> cross-checks: two-peak FFT of xA and of xB, beat period from the envelope, peak-height ratios,
       window comparison (rect / Hann / Hamming / Blackman), decay time.
For every separation: mean, half-range (IB / international-essay method), FFT resolution 1/T, adopted uncertainty,
    f_avg, f_beat, f2^2 - f1^2, ln values, all with propagated uncertainties.
Fits: ln(f2^2 - f1^2) = ln C - n ln d  (best, max and min gradient lines through the error bars; LINEST errors),
    fixed exponent n = 5, f1 against d.
Outputs (in --out):  results.xlsx (tables + native Excel charts with error bars + embedded graphs),
    trials.csv, graph1..graph5 (.png/.pdf), figs/trial_<d>_<n>.png, tables/*.tex (booktabs, \\input-ready), summary.txt
"""
import argparse, os, re, sys, glob, math, csv, io
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.signal import hilbert, find_peaks
from scipy.optimize import curve_fit

# ============================================================== reading Tracker files
def read_tracker(fn):
    """Return t, xA, xB (xB may be None), names. Handles '#multi:' exports and single-track exports."""
    names, header, rows = [], None, []
    with open(fn, encoding="utf-8", errors="ignore") as fh:
        for line in fh:
            line = line.rstrip("\r\n")
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            parts = line.split("\t") if "\t" in line else re.split(r"[,; ]+", line.strip())
            try:
                rows.append([float(p) if p.strip() else np.nan for p in parts])
            except ValueError:
                toks = [p.strip() for p in parts]
                if toks and toks[0].lower() in ("t", "t (s)", "time"):
                    header = [tk.lower() for tk in toks]
                else:
                    names += [tk for tk in toks if tk]
    if not rows:
        raise ValueError("no numeric rows")
    ncol = max(len(r) for r in rows); arr = np.full((len(rows), ncol), np.nan)
    for i, r in enumerate(rows): arr[i, :len(r)] = r
    t = arr[:, 0]
    xcols = []
    if header:                                     # columns called 'x' (skip 'y', 'vx' ...)
        xcols = [i for i, h in enumerate(header) if 0 < i < ncol and h.strip() and h.split()[0] in ("x", "x (m)", "x(m)", "x(mm)")]
    if not xcols:                                  # no usable header: take every column that has data
        xcols = [i for i in range(1, ncol) if np.isfinite(arr[:, i]).sum() > 0.5 * len(arr)]
    if not xcols: raise ValueError("no x columns found")
    xA = arr[:, xcols[0]]; xB = arr[:, xcols[1]] if len(xcols) > 1 else None
    return t, xA, xB, names

def motion_start(t, x, fs):
    """Index where a mass starts to move (end of a flat hold segment, if any)."""
    n0 = max(3, int(0.1 * fs)); hold = np.median(x[:n0]); rng = 0.5 * (np.nanmax(x) - np.nanmin(x))
    if not np.isfinite(rng) or rng <= 0: return 0
    beyond = np.where(np.abs(x - hold) > 0.15 * rng)[0]
    if len(beyond) == 0: return 0
    j = beyond[0]
    while j > 0 and abs(x[j] - hold) > 0.05 * rng: j -= 1
    return j

# ============================================================== spectra
def window(n, name):
    return {"hann": np.hanning, "hamming": np.hamming, "blackman": np.blackman}.get(name, np.ones)(n)

def spectrum(x, fs, win="rect", nfft=2 ** 18):
    w = window(len(x), win); X = np.abs(np.fft.rfft((x - x.mean()) * w, nfft)) * 2 / w.sum()
    return np.fft.rfftfreq(nfft, 1 / fs), X

def peak_interp(f, X, i):
    """Parabolic refinement of a peak at index i."""
    if 0 < i < len(X) - 1:
        a, b, c = X[i - 1], X[i], X[i + 1]; den = a - 2 * b + c
        if den != 0: return f[i] + 0.5 * (a - c) / den * (f[1] - f[0])
    return f[i]

def one_peak(f, X, fmin, fmax):
    m = np.where((f > fmin) & (f < fmax))[0]; i = m[np.argmax(X[m])]
    return peak_interp(f, X, i), X[i]

def two_peaks(f, X, fmin, fmax, T):
    """Two largest local maxima at least one bin (1/T) apart. Returns [(f1,h1),(f2,h2)] sorted by f, or one."""
    m = np.where((f > fmin) & (f < fmax))[0]; df = f[1] - f[0]
    pk, _ = find_peaks(X[m], distance=max(1, int(0.8 / T / df)), prominence=0.05 * X[m].max())
    pk = m[pk]; pk = pk[np.argsort(X[pk])[::-1]][:2]
    return sorted([(peak_interp(f, X, i), X[i]) for i in pk])

# ============================================================== one trial
def analyse_trial(fn, d_mm, args):
    t, xA, xB, names = read_tracker(fn)
    ok = np.isfinite(t) & np.isfinite(xA) & (np.isfinite(xB) if xB is not None else True)
    t, xA = t[ok], xA[ok]; xB = xB[ok] if xB is not None else None
    if args.unit == "m": xA = xA * 1e3; xB = xB * 1e3 if xB is not None else None       # -> mm
    if len(t) < 50: raise ValueError("too few samples")
    dt = (t[-1] - t[0]) / (len(t) - 1); fs = 1 / dt                                    # exact mean spacing (Tracker rounds t to 1 ms)
    res = dict(file=os.path.basename(fn), d_mm=d_mm, fs=fs, N_raw=len(t), t_end_raw=t[-1])
    # ---- release detection: start when BOTH masses are moving
    sA = motion_start(t, xA, fs); sB = motion_start(t, xB, fs) if xB is not None else 0
    j0 = max(sA, sB); t0 = t[j0] + args.skip
    pre = np.nan
    if xB is not None and sA - sB > int(0.05 * fs):          # B already moving while A is still held
        seg = xB[sB:sA]; pre = 0.5 * (seg.max() - seg.min())
    res.update(t_start_A=t[sA], t_start_B=t[sB] if xB is not None else np.nan, t_start=t0, B_preswing_mm=pre)
    keep = t >= t0 - 1e-9; t, xA = t[keep], xA[keep]; xB = xB[keep] if xB is not None else None
    # uniform grid (drop-frame safe)
    tu = np.arange(t[0], t[-1] + dt / 2, dt); xA = np.interp(tu, t, xA); xB = np.interp(tu, t, xB) if xB is not None else None; t = tu
    T = t[-1] - t[0]; res.update(T=T, N=len(t), bin_Hz=1 / T, halfbin_Hz=0.5 / T)
    res["d_track_mm"] = (xB.mean() - xA.mean()) if xB is not None else np.nan          # should equal d if the same feature is tracked on both magnets
    xA = xA - xA.mean(); xB = xB - xB.mean() if xB is not None else None
    res["A_mm"] = 0.5 * (np.percentile(xA[: int(2 * fs)], 99.5) - np.percentile(xA[: int(2 * fs)], 0.5))   # initial amplitude of A (approx)
    # ---- normal coordinates
    if xB is not None:
        eta, xi = xA + xB, xB - xA
        fe, Xe = spectrum(eta, fs, args.window); fx, Xx = spectrum(xi, fs, args.window)
        f1, h1 = one_peak(fe, Xe, args.fmin, args.fmax); f2, h2 = one_peak(fx, Xx, args.fmin, args.fmax)
        # purity: how much of the OTHER mode leaks into each normal coordinate (0 = perfect)
        res["eta_leak"] = np.interp(f2, fe, Xe) / h1; res["xi_leak"] = np.interp(f1, fx, Xx) / h2
        method = "eta/xi"
    else:
        fa, Xa = spectrum(xA, fs, "hann"); pk = two_peaks(fa, Xa, args.fmin, args.fmax, T)
        f1, f2 = (pk[0][0], pk[1][0]) if len(pk) == 2 else (pk[0][0], np.nan); method = "two-peak(A)"
        res["eta_leak"] = res["xi_leak"] = np.nan
    res.update(method=method, f1=f1, f2=f2)
    # ---- window spread (uncertainty from the choice of window)
    WINDOWS = ("rect", "hann", "hamming", "blackman")
    if xB is not None:
        w1, w2 = [], []
        for wn in WINDOWS:
            fe_, Xe_ = spectrum(eta, fs, wn); fx_, Xx_ = spectrum(xi, fs, wn)
            w1.append(one_peak(fe_, Xe_, args.fmin, args.fmax)[0]); w2.append(one_peak(fx_, Xx_, args.fmin, args.fmax)[0])
            res[f"f1_{wn}"], res[f"f2_{wn}"] = w1[-1], w2[-1]
        res["f1_winspread"] = (max(w1) - min(w1)) / 2; res["f2_winspread"] = (max(w2) - min(w2)) / 2
    else:
        for wn in WINDOWS:
            fa_, Xa_ = spectrum(xA, fs, wn); res[f"f1_{wn}"] = one_peak(fa_, Xa_, args.fmin, args.fmax)[0]; res[f"f2_{wn}"] = np.nan
        res["f1_winspread"] = res["f2_winspread"] = np.nan
    for wn in WINDOWS:                                   # conventional method: two peaks in the FFT of mass A alone
        fa_, Xa_ = spectrum(xA, fs, wn); pk_ = two_peaks(fa_, Xa_, args.fmin, args.fmax, T)
        res[f"f1A_{wn}"], res[f"f2A_{wn}"] = (pk_[0][0], pk_[1][0]) if len(pk_) == 2 else (pk_[0][0], np.nan)
    # ---- cross-checks on the single records (Hann, two peaks)
    fa, Xa = spectrum(xA, fs, "hann"); pkA = two_peaks(fa, Xa, args.fmin, args.fmax, T)
    res["f1_A"], res["f2_A"] = (pkA[0][0], pkA[1][0]) if len(pkA) == 2 else (pkA[0][0], np.nan)
    res["rA"] = np.interp(f2, fa, Xa) / np.interp(f1, fa, Xa) if np.isfinite(f2) else np.nan
    if xB is not None:
        fb, Xb = spectrum(xB, fs, "hann"); pkB = two_peaks(fb, Xb, args.fmin, args.fmax, T)
        res["f1_B"], res["f2_B"] = (pkB[0][0], pkB[1][0]) if len(pkB) == 2 else (pkB[0][0], np.nan)
        res["rB"] = np.interp(f2, fb, Xb) / np.interp(f1, fb, Xb) if np.isfinite(f2) else np.nan
    else:
        res["f1_B"] = res["f2_B"] = res["rB"] = np.nan
    res["ratio_pred"] = (f1 / f2) ** 2 if np.isfinite(f2) else np.nan
    # ---- envelope: beat period (time domain) and decay time
    env = np.abs(hilbert(xA)); k = max(1, int(0.2 * fs / f1)); env_s = np.convolve(env, np.ones(k) / k, "same")
    edge = int(fs / f1); env_s[:edge] = np.nan; env_s[-edge:] = np.nan
    Tb_fft = 1 / (f2 - f1) if np.isfinite(f2) and f2 > f1 else np.nan
    res.update(T_beat_fft=Tb_fft, T_beat_env=np.nan, dT_beat_env=np.nan, n_beats=np.nan, tau_s=np.nan, env_minmax=np.nan)
    if np.isfinite(Tb_fft) and T > 1.5 * Tb_fft:
        hi = np.nanmax(env_s)
        imax, _ = find_peaks(np.nan_to_num(env_s, nan=0), distance=max(1, int(0.6 * Tb_fft * fs)), height=0.3 * hi, prominence=0.1 * hi)
        imin, _ = find_peaks(-np.nan_to_num(env_s, nan=hi), distance=max(1, int(0.6 * Tb_fft * fs)), prominence=0.1 * hi)
        if len(imax) >= 3:
            p = np.polyfit(np.arange(len(imax)), t[imax], 1, cov=True); res["T_beat_env"] = p[0][0]; res["dT_beat_env"] = math.sqrt(p[1][0, 0]); res["n_beats"] = len(imax) - 1
            try:
                pe, _ = curve_fit(lambda tt, a, tau: a * np.exp(-tt / tau), t[imax] - t[imax][0], env_s[imax], p0=[env_s[imax][0], 20]); res["tau_s"] = pe[1]
            except Exception: pass
        if len(imax) and len(imin): res["env_minmax"] = np.median(env_s[imin]) / np.median(env_s[imax])
    elif len(t) > 10:                                  # single spring: decay time from the envelope directly
        try:
            m = np.isfinite(env_s); pe, _ = curve_fit(lambda tt, a, tau: a * np.exp(-tt / tau), t[m] - t[m][0], env_s[m], p0=[np.nanmax(env_s), 20]); res["tau_s"] = pe[1]
        except Exception: pass
    # ---- per-trial figure
    if not args.no_trial_figs:
        fig, ax = plt.subplots(2, 2, figsize=(12, 6.2))
        span = min(T, 3.2 * Tb_fft if np.isfinite(Tb_fft) else 3.0); m = t - t[0] <= span
        ax[0, 0].plot(t[m], xA[m], lw=0.8, color="#b35900", label=names[0] if names else "mass A")
        if xB is not None: ax[0, 0].plot(t[m], xB[m], lw=0.8, color="tab:blue", label=names[1] if len(names) > 1 else "mass B")
        ax[0, 0].set_title(f"{res['file']}: first {span:.1f} s after release (t0 = {t0:.2f} s)", fontsize=9); ax[0, 0].legend(fontsize=8); ax[0, 0].set_ylabel("x / mm"); ax[0, 0].set_xlabel("t / s")
        ax[0, 1].plot(t, xA, lw=0.4, color="#b35900"); ax[0, 1].plot(t, env_s, "--", color="crimson", lw=1); ax[0, 1].plot(t, -env_s, "--", color="crimson", lw=1)
        ax[0, 1].set_title(f"mass A, full record T = {T:.1f} s; envelope: T_beat = {res['T_beat_env']:.3f} s, tau = {res['tau_s']:.0f} s", fontsize=9); ax[0, 1].set_xlabel("t / s")
        mm = (fa > f1 - 1.2) & (fa < (f2 if np.isfinite(f2) else f1) + 1.2)
        ax[1, 0].plot(fa[mm], Xa[mm] / Xa[mm].max(), color="#b35900", lw=1, label="|FFT| mass A (Hann)")
        if xB is not None: ax[1, 0].plot(fb[mm], Xb[mm] / Xb[mm].max(), color="tab:blue", lw=1, alpha=0.8, label="|FFT| mass B (Hann)")
        for fq, lab in ((f1, "f1"), (f2, "f2")):
            if np.isfinite(fq): ax[1, 0].axvline(fq, color="0.5", ls=":", lw=0.8); ax[1, 0].text(fq, 1.03, f"{lab} = {fq:.3f} Hz", ha="center", fontsize=8)
        ax[1, 0].set_ylim(0, 1.15); ax[1, 0].set_xlabel("frequency / Hz"); ax[1, 0].set_ylabel("normalised"); ax[1, 0].legend(fontsize=8, loc="upper right")
        ax[1, 0].set_title(f"peak-height ratio: A {res['rA']:.2f}, B {res['rB']:.2f}; (f1/f2)^2 = {res['ratio_pred']:.2f}", fontsize=9)
        if xB is not None:
            ax[1, 1].semilogy(fe[mm], Xe[mm] / max(Xe[mm].max(), 1e-12), color="tab:green", lw=1, label="|FFT| of eta = xA + xB")
            ax[1, 1].semilogy(fx[mm], Xx[mm] / max(Xx[mm].max(), 1e-12), color="tab:purple", lw=1, label="|FFT| of xi = xB - xA")
            ax[1, 1].set_ylim(1e-3, 2); ax[1, 1].legend(fontsize=8, loc="upper right"); ax[1, 1].set_xlabel("frequency / Hz")
            ax[1, 1].set_title(f"normal coordinates ({args.window}): f1 = {f1:.4f}, f2 = {f2:.4f} Hz; leakage {res['eta_leak']:.3f}/{res['xi_leak']:.3f}", fontsize=9)
        else:
            ax[1, 1].axis("off")
        for a in ax.flat: a.grid(alpha=0.3)
        fig.tight_layout(); os.makedirs(os.path.join(args.out, "figs"), exist_ok=True)
        fig.savefig(os.path.join(args.out, "figs", "trial_" + os.path.splitext(res["file"])[0] + ".png"), dpi=130); plt.close(fig)
    return res

# ============================================================== helpers
TEX = False
def fmt_pm(v, u, unit=""):
    """IB style: uncertainty to 1 s.f. (2 if it starts with 1), value to the same decimal place."""
    if not (np.isfinite(v) and np.isfinite(u)) or u <= 0: return f"{v:.4g}{unit}"
    sf = 2 if f"{u:.1e}"[0] == "1" else 1
    dec = max(0, -int(math.floor(math.log10(u))) + sf - 1)
    return f"{v:.{dec}f} $\\pm$ {u:.{dec}f}{unit}" if TEX else f"{v:.{dec}f} ± {u:.{dec}f}{unit}"

def half_range(a): a = np.asarray(a, float); a = a[np.isfinite(a)]; return (a.max() - a.min()) / 2 if len(a) > 1 else np.nan

def parse_d(name):
    s = name.strip().lower().replace("mm", "").replace(",", ".")
    if s in ("inf", "infinity", "single", "∞", "far"): return math.inf
    m = re.match(r"^\s*([0-9]+(?:\.[0-9]+)?)", s); return float(m.group(1)) if m else None

# ============================================================== main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder"); ap.add_argument("--out", default="results")
    ap.add_argument("--unit", default="m", choices=["m", "mm"], help="unit of x in the Tracker files")
    ap.add_argument("--dd", type=float, default=0.5, help="uncertainty of d in mm (default 0.5)")
    ap.add_argument("--unc", default="max", choices=["max", "halfrange", "halfbin", "sd"], help="per-separation uncertainty of f1, f2: max(half-range, half-bin) [default], half-range only, half-bin only, or standard deviation")
    ap.add_argument("--window", default="rect", choices=["rect", "hann", "hamming", "blackman"], help="window for the eta/xi peak positions")
    ap.add_argument("--fmin", type=float, default=3.0); ap.add_argument("--fmax", type=float, default=15.0)
    ap.add_argument("--skip", type=float, default=0.0, help="extra seconds discarded after the detected release")
    ap.add_argument("--no-trial-figs", action="store_true")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True); log = io.StringIO()
    def say(*a):
        s = " ".join(str(x) for x in a); print(s); log.write(s + "\n")

    # ---------- collect files
    items = []
    for root, dirs, files in os.walk(args.folder):
        for fn in sorted(files):
            if not fn.lower().endswith((".txt", ".csv", ".dat")): continue
            d = parse_d(os.path.basename(root)) if os.path.abspath(root) != os.path.abspath(args.folder) else None
            if d is None: d = parse_d(fn.split("_")[0])
            if d is None: say(f"skip {fn}: cannot read d from folder or file name"); continue
            m = re.search(r"_(\d+)", os.path.splitext(fn)[0]); trial = int(m.group(1)) if m else len([i for i in items if i[0] == d]) + 1
            items.append((d, trial, os.path.join(root, fn)))
    items.sort(key=lambda x: (x[0], x[1]))
    if not items: sys.exit("no files found")
    say(f"{len(items)} files, separations: {sorted(set(i[0] for i in items))}")

    # ---------- per-trial analysis
    trials = []
    for d, trial, fn in items:
        try:
            r = {"trial": trial}; r.update(analyse_trial(fn, d, args)); r["trial"] = trial
        except Exception as e:
            say(f"  !! {fn}: {e}"); continue
        trials.append(r)
        warn = []
        if np.isfinite(r["d_track_mm"]) and np.isfinite(d) and abs(r["d_track_mm"] - d) > 1.0: warn.append(f"mean(xB) - mean(xA) = {r['d_track_mm']:.1f} mm but folder says d = {d} mm: if both templates are centred on the magnets this IS d (check ruler value / calibration stick); if the templates sit on different features the offset is harmless")
        if np.isfinite(r.get("B_preswing_mm", np.nan)) and r["B_preswing_mm"] > 0.3:
            warn.append(f"mass B was already swinging (+-{r['B_preswing_mm']:.1f} mm) during {r['t_start_B']:.2f}-{r['t_start_A']:.2f} s while A was still held -> f1, f2 unaffected (analysis starts at A's release), but the peak-height-ratio / beat-minimum tests do not apply to this trial (hold A, pull it slowly, wait until B is still, release)")
        if abs(r["fs"] / round(r["fs"]) - 1) > 0.002 and abs(r["fs"] / (round(r["fs"] / 29.97) * 29.97) - 1) > 0.002: warn.append(f"fs = {r['fs']:.2f} Hz is not a round frame rate: fine if this is the true rate from the video-file properties typed into Clip Settings (a 0.5 % error in fs shifts every frequency by 0.5 %)")
        if np.isfinite(r.get("T_beat_fft", np.nan)) and r["T"] < 3 * r["T_beat_fft"]: warn.append(f"record {r['T']:.1f} s < 3 T_beat = {3*r['T_beat_fft']:.1f} s")
        say(f"d = {d} mm trial {trial}: T = {r['T']:.1f} s, fs = {r['fs']:.2f} Hz, f1 = {r['f1']:.4f}, f2 = {r['f2']:.4f} Hz ({r['method']}), "
            f"A-only: {r['f1_A']:.3f}/{r['f2_A']:.3f}, T_beat env {r['T_beat_env']:.3f} vs 1/(f2-f1) {r['T_beat_fft']:.3f} s, rA {r['rA']:.2f} rB {r['rB']:.2f} pred {r['ratio_pred']:.2f}"
            + ("".join("\n      WARNING: " + w for w in warn)))
    if not trials: sys.exit("nothing analysed")
    cols = list(trials[0].keys())
    with open(os.path.join(args.out, "trials.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols); w.writeheader(); [w.writerow(r) for r in trials]

    # ---------- per-separation aggregation
    ds = sorted(set(r["d_mm"] for r in trials if np.isfinite(r["d_mm"]))); agg = []
    for d in ds:
        rs = [r for r in trials if r["d_mm"] == d]; f1s = np.array([r["f1"] for r in rs]); f2s = np.array([r["f2"] for r in rs])
        hb = np.mean([r["halfbin_Hz"] for r in rs]); ws1 = np.nanmean([r["f1_winspread"] for r in rs]); ws2 = np.nanmean([r["f2_winspread"] for r in rs])
        def unc(vals):
            hr, sd = half_range(vals), (np.nanstd(vals, ddof=1) if len(vals) > 1 else np.nan)
            return {"max": np.nanmax([hr, hb]), "halfrange": hr, "halfbin": hb, "sd": sd}[args.unc]
        a = dict(d_mm=d, dd_mm=args.dd, n_trials=len(rs), T_mean=np.mean([r["T"] for r in rs]), halfbin=hb,
                 f1=np.nanmean(f1s), f1_hr=half_range(f1s), f1_sd=np.nanstd(f1s, ddof=1) if len(rs) > 1 else np.nan, f1_ws=ws1, f1_u=unc(f1s),
                 f2=np.nanmean(f2s), f2_hr=half_range(f2s), f2_sd=np.nanstd(f2s, ddof=1) if len(rs) > 1 else np.nan, f2_ws=ws2, f2_u=unc(f2s),
                 f1_trials=list(f1s), f2_trials=list(f2s), rA=np.nanmean([r["rA"] for r in rs]), rB=np.nanmean([r["rB"] for r in rs]),
                 tau=np.nanmean([r["tau_s"] for r in rs]), d_track=np.nanmean([r["d_track_mm"] for r in rs]), env_minmax=np.nanmean([r["env_minmax"] for r in rs]))
        a["f_avg"] = (a["f1"] + a["f2"]) / 2; a["f_avg_u"] = (a["f1_u"] + a["f2_u"]) / 2
        a["f_beat"] = a["f2"] - a["f1"]; a["f_beat_u"] = a["f1_u"] + a["f2_u"]
        a["y"] = a["f2"] ** 2 - a["f1"] ** 2; a["y_u"] = 2 * a["f2"] * a["f2_u"] + 2 * a["f1"] * a["f1_u"]
        a["ln_d"] = math.log(d / 1e3); a["ln_d_u"] = args.dd / d
        a["ln_y"] = math.log(a["y"]) if a["y"] > 0 else np.nan; a["ln_y_u"] = a["y_u"] / a["y"] if a["y"] > 0 else np.nan
        a["C_i"] = a["y"] * (d / 1e3) ** 5; a["C_i_u"] = a["C_i"] * (a["y_u"] / a["y"] + 5 * args.dd / d) if a["y"] > 0 else np.nan
        a["ratio_pred"] = (a["f1"] / a["f2"]) ** 2; a["coupling"] = (a["f2"] / a["f1"]) ** 2 - 1
        for wn in ("rect", "hann", "hamming", "blackman"):                  # window comparison (Graph 6)
            v1 = np.array([r.get(f"f1_{wn}", np.nan) for r in rs]); v2 = np.array([r.get(f"f2_{wn}", np.nan) for r in rs])
            a[f"f1_{wn}"], a[f"f2_{wn}"] = np.nanmean(v1), np.nanmean(v2); a[f"f1_{wn}_u"], a[f"f2_{wn}_u"] = unc(v1), unc(v2)
            a[f"y_{wn}"] = a[f"f2_{wn}"] ** 2 - a[f"f1_{wn}"] ** 2; a[f"y_{wn}_u"] = 2 * a[f"f2_{wn}"] * a[f"f2_{wn}_u"] + 2 * a[f"f1_{wn}"] * a[f"f1_{wn}_u"]
            v1 = np.array([r.get(f"f1A_{wn}", np.nan) for r in rs]); v2 = np.array([r.get(f"f2A_{wn}", np.nan) for r in rs])
            a[f"f1A_{wn}"], a[f"f2A_{wn}"] = np.nanmean(v1), np.nanmean(v2); a[f"f1A_{wn}_u"], a[f"f2A_{wn}_u"] = unc(v1), unc(v2)
            a[f"yA_{wn}"] = a[f"f2A_{wn}"] ** 2 - a[f"f1A_{wn}"] ** 2; a[f"yA_{wn}_u"] = 2 * a[f"f2A_{wn}"] * a[f"f2A_{wn}_u"] + 2 * a[f"f1A_{wn}"] * a[f"f1A_{wn}_u"]
        agg.append(a)
    inf_rows = [r for r in trials if not np.isfinite(r["d_mm"])]

    # ---------- fits
    fit = {}
    good = [a for a in agg if np.isfinite(a["ln_y"])]
    if len(good) >= 3:
        X = np.array([a["ln_d"] for a in good]); Y = np.array([a["ln_y"] for a in good]); dX = np.array([a["ln_d_u"] for a in good]); dY = np.array([a["ln_y_u"] for a in good])
        p, cov = np.polyfit(X, Y, 1, cov=True); n_best, lnC = -p[0], p[1]; se_n, se_lnC = math.sqrt(cov[0, 0]), math.sqrt(cov[1, 1])
        Yhat = np.polyval(p, X); R2 = 1 - np.sum((Y - Yhat) ** 2) / np.sum((Y - Y.mean()) ** 2)
        i0, i1 = np.argmin(X), np.argmax(X)                                                   # smallest d (largest y) and largest d
        steep = ((Y[i1] - dY[i1]) - (Y[i0] + dY[i0])) / ((X[i1] + dX[i1]) - (X[i0] - dX[i0]))    # most negative gradient
        shallow = ((Y[i1] + dY[i1]) - (Y[i0] - dY[i0])) / ((X[i1] - dX[i1]) - (X[i0] + dX[i0]))  # least negative gradient
        c_steep = (Y[i0] + dY[i0]) - steep * (X[i0] - dX[i0]); c_shallow = (Y[i0] - dY[i0]) - shallow * (X[i0] + dX[i0])
        fit.update(n_best=n_best, se_n=se_n, lnC=lnC, se_lnC=se_lnC, R2=R2, n_max=-steep, n_min=-shallow, dn_maxmin=(abs(steep) - abs(shallow)) / 2,
                   c_steep=c_steep, c_shallow=c_shallow, C=math.exp(lnC), C_u_maxmin=math.exp(lnC) * abs(c_steep - c_shallow) / 2, C_u_se=math.exp(lnC) * se_lnC,
                   X=X, Y=Y, dX=dX, dY=dY, resid=Y - Yhat)
        # fixed exponent n = 5
        dm = np.array([a["d_mm"] for a in good]) / 1e3; y = np.array([a["y"] for a in good]); yu = np.array([a["y_u"] for a in good])
        C5 = np.sum(y * dm ** -5) / np.sum(dm ** -10); Ci = y * dm ** 5
        fit.update(C5_ls=C5, C5_mean=Ci.mean(), C5_hr=half_range(Ci), C5_sd=Ci.std(ddof=1) if len(Ci) > 1 else np.nan,
                   chi2_n5=np.sum(((y - C5 * dm ** -5) / yu) ** 2), dof_n5=len(y) - 1)
        # f1 against d
        f1v = np.array([a["f1"] for a in agg]); f1u = np.array([a["f1_u"] for a in agg]); dd_ = np.array([a["d_mm"] for a in agg])
        pf, covf = np.polyfit(dd_, f1v, 1, cov=True) if len(agg) > 2 else ((np.nan, np.nan), np.full((2, 2), np.nan))
        fit.update(f1_mean=f1v.mean(), f1_hr=half_range(f1v), f1_sd=f1v.std(ddof=1) if len(f1v) > 1 else np.nan, f1_slope=pf[0], f1_slope_se=math.sqrt(covf[0, 0]) if np.isfinite(covf[0, 0]) else np.nan,
                   f1_max_dev=np.max(np.abs(f1v - f1v.mean())), f1_mean_unc=np.mean(f1u))
        say("\nFITS")
        say(f"  ln-ln: n = {n_best:.3f} (LINEST s.e. {se_n:.3f}); max/min gradient lines: n_max = {-steep:.3f}, n_min = {-shallow:.3f} -> n = {fmt_pm(n_best, fit['dn_maxmin'])}; R^2 = {R2:.4f}")
        say(f"         ln C = {lnC:.3f} +- {se_lnC:.3f} -> C = {fit['C']:.3e} m^5 s^-2 (+- {fit['C_u_maxmin']:.1e} from max/min lines)")
        say(f"  fixed n = 5: C = {C5:.3e} (least squares); mean of C_i = {fmt_pm(Ci.mean(), half_range(Ci))}; chi^2/dof = {fit['chi2_n5']:.1f}/{fit['dof_n5']}")
        say(f"  f1: mean {f1v.mean():.4f} Hz, half-range {half_range(f1v):.4f}, slope vs d {pf[0]*1e3:.2f} +- {fit['f1_slope_se']*1e3:.2f} mHz/mm, largest deviation {fit['f1_max_dev']:.4f} Hz vs mean uncertainty {np.mean(f1u):.4f} Hz")
    if inf_rows:
        say("  Infinity (uncoupled) recordings: " + ", ".join(f"{r['file']}: fA = {r['f1_A']:.4f}" + (f", fB = {r['f1_B']:.4f}" if np.isfinite(r['f1_B']) else "") + f" Hz (T = {r['T']:.0f} s)" for r in inf_rows))

    # ---------- summary graphs (matplotlib: both error bars, max/min lines)
    dmm = np.array([a["d_mm"] for a in agg]); C1, C2 = "#b35900", "crimson"
    if fit:
        dsm = np.linspace(0.9 * dmm.min(), 1.08 * dmm.max(), 200) / 1e3; f1m = fit["f1_mean"]
        # graph 1
        fig, ax = plt.subplots(figsize=(7.2, 4.4))
        ax.errorbar(dmm, [a["f2"] for a in agg], yerr=[a["f2_u"] for a in agg], xerr=args.dd, fmt="s", color=C2, ms=4, capsize=2, lw=1, label="$f_2$ (anti-phase), measured")
        ax.errorbar(dmm, [a["f1"] for a in agg], yerr=[a["f1_u"] for a in agg], xerr=args.dd, fmt="o", color=C1, ms=4, capsize=2, lw=1, label="$f_1$ (in-phase), measured")
        ax.plot(dsm * 1e3, np.sqrt(f1m ** 2 + fit["C5_ls"] * dsm ** -5), "-", color=C2, lw=1.2, label=rf"$\sqrt{{f_1^2 + C d^{{-5}}}}$, $C$ fitted ($n=5$)")
        ax.plot(dsm * 1e3, np.sqrt(f1m ** 2 + fit["C"] * dsm ** -fit["n_best"]), "--", color="0.4", lw=1, label=rf"free exponent fit, $n={fit['n_best']:.2f}$")
        ax.axhline(f1m, color=C1, lw=1, ls=":", label=rf"$\bar f_1={f1m:.3f}$ Hz")
        ax.set_xlabel("centre-to-centre separation $d$ / mm"); ax.set_ylabel("frequency / Hz"); ax.grid(alpha=0.3); ax.legend(fontsize=8)
        ax.set_title("Graph 1: normal-mode frequencies against separation", fontsize=10); fig.tight_layout()
        fig.savefig(os.path.join(args.out, "graph1_frequencies.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph1_frequencies.pdf")); plt.close(fig)
        # graph 2: ln-ln
        X, Y, dX, dY = fit["X"], fit["Y"], fit["dX"], fit["dY"]; xx = np.array([X.min() - 1.5 * dX.max(), X.max() + 1.5 * dX.max()])
        fig, ax = plt.subplots(figsize=(7.2, 4.6))
        ax.errorbar(X, Y, yerr=dY, xerr=dX, fmt="o", color="k", ms=4, capsize=2, lw=1, label="data")
        ax.plot(xx, fit["lnC"] - fit["n_best"] * xx, "-", color=C2, lw=1.3, label=rf"best fit: gradient $={-fit['n_best']:.2f}$, $R^2={fit['R2']:.4f}$")
        ax.plot(xx, fit["c_steep"] - fit["n_max"] * xx, "--", color="0.4", lw=1, label=rf"max gradient $={-fit['n_max']:.2f}$")
        ax.plot(xx, fit["c_shallow"] - fit["n_min"] * xx, ":", color="0.4", lw=1.2, label=rf"min gradient $={-fit['n_min']:.2f}$")
        ax.plot(xx, Y.mean() - 5 * (xx - X.mean()), "-", color="tab:blue", lw=0.8, alpha=0.6, label="gradient $-5$ (dipole model) through the centroid")
        ax.set_xlabel(r"$\ln(d\,/\,\mathrm{m})$"); ax.set_ylabel(r"$\ln\left[(f_2^2-f_1^2)\,/\,\mathrm{Hz^2}\right]$"); ax.grid(alpha=0.3); ax.legend(fontsize=8)
        ax.set_title(rf"Graph 2: $n = {fit['n_best']:.2f}\pm{fit['dn_maxmin']:.2f}$ (max/min gradient method)", fontsize=10); fig.tight_layout()
        fig.savefig(os.path.join(args.out, "graph2_loglog.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph2_loglog.pdf")); plt.close(fig)
        # graph 3: f_beat and f_avg
        fig, ax = plt.subplots(1, 2, figsize=(10, 4))
        ax[0].errorbar(dmm, [a["f_beat"] for a in agg], yerr=[a["f_beat_u"] for a in agg], xerr=args.dd, fmt="o", color=C2, ms=4, capsize=2, lw=1, label="measured $f_2-f_1$")
        ax[0].plot(dsm * 1e3, np.sqrt(f1m ** 2 + fit["C5_ls"] * dsm ** -5) - f1m, "-", color=C2, lw=1.2, label="dipole model, $C$ fitted")
        ax[0].set_xlabel("$d$ / mm"); ax[0].set_ylabel("$f_{beat}=f_2-f_1$ / Hz"); ax[0].set_yscale("log"); ax[0].grid(alpha=0.3, which="both"); ax[0].legend(fontsize=8); ax[0].set_title("Graph 3a: beat frequency", fontsize=10)
        ax[1].errorbar(dmm, [a["f_avg"] for a in agg], yerr=[a["f_avg_u"] for a in agg], xerr=args.dd, fmt="o", color="tab:blue", ms=4, capsize=2, lw=1, label="measured $(f_1+f_2)/2$")
        ax[1].plot(dsm * 1e3, 0.5 * (f1m + np.sqrt(f1m ** 2 + fit["C5_ls"] * dsm ** -5)), "-", color="tab:blue", lw=1.2, label="dipole model")
        ax[1].set_xlabel("$d$ / mm"); ax[1].set_ylabel("$f_{avg}$ / Hz"); ax[1].grid(alpha=0.3); ax[1].legend(fontsize=8); ax[1].set_title("Graph 3b: average frequency", fontsize=10)
        fig.tight_layout(); fig.savefig(os.path.join(args.out, "graph3_beat_avg.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph3_beat_avg.pdf")); plt.close(fig)
        # graph 4: C_i = (f2^2-f1^2) d^5 against d (should be flat) + residuals of the ln-ln fit
        fig, ax = plt.subplots(1, 2, figsize=(10, 4))
        ax[0].errorbar(dmm, [a["C_i"] for a in agg], yerr=[a["C_i_u"] for a in agg], xerr=args.dd, fmt="o", color="k", ms=4, capsize=2, lw=1)
        ax[0].axhline(fit["C5_ls"], color=C2, lw=1.2, label=rf"$C={fit['C5_ls']:.2e}$ m$^5$s$^{{-2}}$ (fixed $n=5$)"); ax[0].set_xlabel("$d$ / mm"); ax[0].set_ylabel(r"$(f_2^2-f_1^2)\,d^5$ / m$^5$ s$^{-2}$"); ax[0].grid(alpha=0.3); ax[0].legend(fontsize=8)
        ax[0].set_title("Graph 4a: test of the exponent (flat = $d^{-5}$)", fontsize=10)
        ax[1].errorbar(X, fit["resid"], yerr=dY, fmt="o", color="k", ms=4, capsize=2, lw=1); ax[1].axhline(0, color=C2, lw=1); ax[1].set_xlabel(r"$\ln(d/\mathrm{m})$"); ax[1].set_ylabel("residual of ln–ln fit"); ax[1].grid(alpha=0.3)
        ax[1].set_title("Graph 4b: residuals (a curve = deviation from a power law)", fontsize=10)
        fig.tight_layout(); fig.savefig(os.path.join(args.out, "graph4_exponent_test.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph4_exponent_test.pdf")); plt.close(fig)
        # graph 5: peak-height ratios
        fig, ax = plt.subplots(figsize=(7.2, 4))
        ax.plot(dmm, [a["rA"] for a in agg], "o", color=C1, label="mass A: $|X(f_2)|/|X(f_1)|$"); ax.plot(dmm, [a["rB"] for a in agg], "s", color="tab:blue", label="mass B")
        ax.plot(dmm, [a["ratio_pred"] for a in agg], "-", color="k", lw=1, label="hold-and-release prediction $(f_1/f_2)^2$"); ax.axhline(1, color="0.5", ls=":", lw=1, label="instantaneous release: 1")
        ax.set_xlabel("$d$ / mm"); ax.set_ylabel("FFT peak-height ratio"); ax.grid(alpha=0.3); ax.legend(fontsize=8); ax.set_title("Graph 5: mode amplitudes (initial-condition test)", fontsize=10)
        fig.tight_layout(); fig.savefig(os.path.join(args.out, "graph5_peak_ratio.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph5_peak_ratio.pdf")); plt.close(fig)

    # ---------- window comparison: ln-ln fit repeated with each window and each extraction method (Graph 6, Table)
    winfit = {}
    if fit:
        WN = ("rect", "hann", "hamming", "blackman"); WLAB = dict(rect="rectangular", hann="Hann", hamming="Hamming", blackman="Blackman")
        METH = (("", "normal coordinates: single peak of FFT(xA+xB) and FFT(xB-xA)"), ("A", "conventional: two peaks in FFT(xA) alone"))
        def lnfit(ykey, ukey):
            g = [a for a in agg if np.isfinite(a[ykey]) and a[ykey] > 0]
            Xw = np.array([a["ln_d"] for a in g]); Yw = np.log([a[ykey] for a in g]); dYw = np.array([a[ukey] / a[ykey] for a in g]); dXw = np.array([a["ln_d_u"] for a in g])
            pw, covw = np.polyfit(Xw, Yw, 1, cov=True); i0, i1 = np.argmin(Xw), np.argmax(Xw)
            st = ((Yw[i1] - dYw[i1]) - (Yw[i0] + dYw[i0])) / ((Xw[i1] + dXw[i1]) - (Xw[i0] - dXw[i0])); sh = ((Yw[i1] + dYw[i1]) - (Yw[i0] - dYw[i0])) / ((Xw[i1] - dXw[i1]) - (Xw[i0] + dXw[i0]))
            return dict(n=-pw[0], se=math.sqrt(covw[0, 0]), dn=(abs(st) - abs(sh)) / 2, lnC=pw[1], C=math.exp(pw[1]), R2=1 - np.sum((Yw - np.polyval(pw, Xw)) ** 2) / np.sum((Yw - Yw.mean()) ** 2), X=Xw, Y=Yw, dY=dYw)
        for m, _ in METH:
            for wn in WN:
                w = lnfit(f"y{m}_{wn}", f"y{m}_{wn}_u"); w["f1"] = np.mean([a[f"f1{m}_{wn}"] for a in agg]); w["f1_hr"] = half_range([a[f"f1{m}_{wn}"] for a in agg]); winfit[m + wn] = w
        say("\nWINDOW COMPARISON (same records, same trimming; only the FFT window / extraction method changes)")
        for m, mlab in METH:
            say(f"  {mlab}")
            for wn in WN:
                w = winfit[m + wn]; say(f"    {WLAB[wn]:12s}: n = {w['n']:.3f} +- {w['dn']:.3f} (max/min), s.e. {w['se']:.3f}, R^2 = {w['R2']:.4f}, C = {w['C']:.3e}, mean f1 = {w['f1']:.4f} Hz")
            nn = [winfit[m + wn]["n"] for wn in WN]; say(f"    spread of n over the four windows: +-{(max(nn)-min(nn))/2:.3f}   (max/min-gradient uncertainty of one fit: +-{winfit[m + 'rect']['dn']:.2f})")
        # ---- Graph 6: 2 rows (methods) x 3 panels
        col = {"rect": "k", "hann": "crimson", "hamming": "tab:blue", "blackman": "tab:green"}; mk = {"rect": "o", "hann": "s", "hamming": "^", "blackman": "D"}; off = {"rect": -0.45, "hann": -0.15, "hamming": 0.15, "blackman": 0.45}
        fig, axs = plt.subplots(2, 3, figsize=(15, 8.6))
        for row, (m, mlab) in enumerate(METH):
            ax = axs[row]
            for wn in WN:
                ax[0].errorbar(dmm + off[wn], [a[f"f2{m}_{wn}"] for a in agg], yerr=[a[f"f2{m}_{wn}_u"] for a in agg], fmt=mk[wn], color=col[wn], ms=4, capsize=2, lw=1, label=f"$f_2$, {WLAB[wn]}")
                ax[0].errorbar(dmm + off[wn], [a[f"f1{m}_{wn}"] for a in agg], yerr=[a[f"f1{m}_{wn}_u"] for a in agg], fmt=mk[wn], color=col[wn], ms=4, capsize=2, lw=1, mfc="white")
                w = winfit[m + wn]; ax[1].errorbar(w["X"] + off[wn] * 0.004, w["Y"], yerr=w["dY"], fmt=mk[wn], color=col[wn], ms=4, capsize=2, lw=1, label=f"{WLAB[wn]}: gradient ${-w['n']:.2f}\\pm{w['dn']:.2f}$")
                xx_ = np.array([w["X"].min() - 0.03, w["X"].max() + 0.03]); ax[1].plot(xx_, w["lnC"] - w["n"] * xx_, "-", color=col[wn], lw=0.9, alpha=0.7)
            ax[0].set_xlabel("$d$ / mm"); ax[0].set_ylabel("frequency / Hz"); ax[0].grid(alpha=0.3); ax[0].legend(fontsize=7, ncol=2); ax[0].set_title(f"({'ad'[row]}) $f_1$ (open) and $f_2$ (filled), {mlab.split(':')[0]}", fontsize=9)
            ax[1].set_xlabel(r"$\ln(d/\mathrm{m})$"); ax[1].set_ylabel(r"$\ln[(f_2^2-f_1^2)/\mathrm{Hz^2}]$"); ax[1].grid(alpha=0.3); ax[1].legend(fontsize=7); ax[1].set_title(f"({'be'[row]}) ln–ln fit with each window", fontsize=9)
            xs = np.arange(len(WN)); ax[2].bar(xs, [winfit[m + wn]["n"] for wn in WN], yerr=[winfit[m + wn]["dn"] for wn in WN], color=[col[wn] for wn in WN], alpha=0.75, capsize=4, width=0.6)
            ax[2].axhline(5, color="0.3", ls="--", lw=1, label="dipole model: $n=5$"); ax[2].set_xticks(xs); ax[2].set_xticklabels([WLAB[wn] for wn in WN], fontsize=8); ax[2].set_ylabel("exponent $n$")
            lo = min(winfit[m + wn]["n"] - winfit[m + wn]["dn"] for wn in WN); hi = max(winfit[m + wn]["n"] + winfit[m + wn]["dn"] for wn in WN); ax[2].set_ylim(min(lo, 5) - 0.3, max(hi, 5) + 0.3)
            for i, wn in enumerate(WN): ax[2].text(i, winfit[m + wn]["n"] + winfit[m + wn]["dn"] + 0.05, f"{winfit[m + wn]['n']:.2f}", ha="center", fontsize=8)
            ax[2].grid(alpha=0.3, axis="y"); ax[2].legend(fontsize=8); ax[2].set_title(f"({'cf'[row]}) exponent from each window (error bar: max/min gradient)", fontsize=9)
        fig.suptitle("Graph 6: effect of the FFT window. Top: normal-coordinate method (one peak per spectrum). Bottom: conventional two-peak FFT of mass A alone.", fontsize=10); fig.tight_layout()
        fig.savefig(os.path.join(args.out, "graph6_windows.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph6_windows.pdf")); plt.close(fig)
        # ---- Graph 7: shift of f2 with the window, per point, against the FFT resolution (half-bin) and the random scatter
        fig, ax = plt.subplots(1, 2, figsize=(11, 4.2))
        hb_arr = np.array([a["halfbin"] for a in agg]) * 1e3; hr_arr = np.array([a["f2_hr"] for a in agg]) * 1e3
        for k, (m, mlab) in enumerate(METH):
            ax[k].fill_between(dmm, -hb_arr, hb_arr, color="0.85", label="± half a frequency bin, $1/(2T)$")
            ax[k].plot(dmm, hr_arr, ":", color="0.4", lw=1, label="half-range of the 5 trials"); ax[k].plot(dmm, -hr_arr, ":", color="0.4", lw=1)
            for wn in ("hann", "hamming", "blackman"):
                ax[k].plot(dmm + off[wn], [1e3 * (a[f"f2{m}_{wn}"] - a[f"f2{m}_rect"]) for a in agg], mk[wn], color=col[wn], ms=5, label=f"{WLAB[wn]} − rectangular")
            ax[k].axhline(0, color="k", lw=0.8); ax[k].set_xlabel("$d$ / mm"); ax[k].set_ylabel("$f_2$(window) − $f_2$(rectangular) / mHz"); ax[k].grid(alpha=0.3); ax[k].legend(fontsize=7, loc="upper right")
            ax[k].set_ylim(-1.6 * hb_arr.max(), 1.6 * hb_arr.max())      # scale to the FFT resolution; an off-scale half-range is an outlier trial, visible in the Raw tables
            ax[k].set_title(f"({'ab'[k]}) {mlab.split(':')[0]}", fontsize=9)
        fig.suptitle("Graph 7: is the window choice significant? Shift of $f_2$ between windows compared with the FFT resolution and the trial-to-trial scatter", fontsize=9.5)
        fig.tight_layout(); fig.savefig(os.path.join(args.out, "graph7_window_shift.png"), dpi=200); fig.savefig(os.path.join(args.out, "graph7_window_shift.pdf")); plt.close(fig)
        # ---- table for the appendix
        wt = [r"\begin{tabular}{@{}lcccc@{}}", r"\toprule", r"method / window & rectangular & Hann & Hamming & Blackman \\", r"\midrule"]
        for m, mlab in METH:
            wt.append(rf"\multicolumn{{5}}{{l}}{{\emph{{{mlab}}}}} \\")
            wt.append("exponent $n$ (max/min) & " + " & ".join(rf"${winfit[m + wn]['n']:.2f}\pm{winfit[m + wn]['dn']:.2f}$" for wn in WN) + r" \\")
            wt.append("$n$ (LINEST s.e.) & " + " & ".join(rf"${winfit[m + wn]['n']:.3f}\pm{winfit[m + wn]['se']:.3f}$" for wn in WN) + r" \\")
            wt.append("$R^2$ & " + " & ".join(f"{winfit[m + wn]['R2']:.4f}" for wn in WN) + r" \\")
            wt.append("$C$ / m$^5$ s$^{-2}$ & " + " & ".join(f"{winfit[m + wn]['C']:.3e}" for wn in WN) + r" \\")
            wt.append(r"mean $f_1$ / Hz & " + " & ".join(f"{winfit[m + wn]['f1']:.4f}" for wn in WN) + r" \\")
            if m == "": wt.append(r"\midrule")
        wt += [r"\bottomrule", r"\end{tabular}"]
        os.makedirs(os.path.join(args.out, "tables"), exist_ok=True); open(os.path.join(args.out, "tables", "windows.tex"), "w").write("\n".join(wt))
        # per-d table
        wt = [r"\begin{tabular}{@{}c" + "cc" * 4 + r"@{}}", r"\toprule", r"$d$ / mm & \multicolumn{2}{c}{rectangular} & \multicolumn{2}{c}{Hann} & \multicolumn{2}{c}{Hamming} & \multicolumn{2}{c}{Blackman} \\",
              r" & $f_1$ / Hz & $f_2$ / Hz" * 4 + r" \\", r"\midrule"]
        for a in agg: wt.append(f"{a['d_mm']:g} & " + " & ".join(f"{a[f'f1_{wn}']:.3f} & {a[f'f2_{wn}']:.3f}" for wn in WN) + r" \\")
        wt += [r"\bottomrule", r"\end{tabular}"]; open(os.path.join(args.out, "tables", "windows_per_d.tex"), "w").write("\n".join(wt))

    # ---------- LaTeX tables
    global TEX; TEX = True
    os.makedirs(os.path.join(args.out, "tables"), exist_ok=True); ntr = max(a["n_trials"] for a in agg)
    def tex_raw(key, label):
        L = [r"\begin{tabular}{@{}c" + "c" * ntr + r"cc@{}}", r"\toprule", rf"$d$ / mm & \multicolumn{{{ntr}}}{{c}}{{{label} / Hz, trials 1--{ntr}}} & $T$ / s & {label} / Hz \\", r"\midrule"]
        for a in agg:
            vals = a[key + "_trials"] + [np.nan] * (ntr - len(a[key + "_trials"]))
            L.append(f"{a['d_mm']:g} $\\pm$ {args.dd:g} & " + " & ".join(f"{v:.3f}" if np.isfinite(v) else "--" for v in vals) + f" & {a['T_mean']:.0f} & {fmt_pm(a[key], a[key + '_u'])} \\\\")
        L += [r"\bottomrule", r"\end{tabular}"]; return "\n".join(L)
    open(os.path.join(args.out, "tables", "raw_f1.tex"), "w").write(tex_raw("f1", "$f_1$")); open(os.path.join(args.out, "tables", "raw_f2.tex"), "w").write(tex_raw("f2", "$f_2$"))
    L = [r"\begin{tabular}{@{}cccccccc@{}}", r"\toprule", r"$d$ / mm & $f_1$ / Hz & $f_2$ / Hz & $f_{\rm avg}$ / Hz & $f_{\rm beat}$ / Hz & $f_2^2-f_1^2$ / Hz$^2$ & $\ln(d/\mathrm{m})$ & $\ln[(f_2^2-f_1^2)/\mathrm{Hz^2}]$ \\", r"\midrule"]
    for a in agg:
        L.append(f"{a['d_mm']:g} $\\pm$ {args.dd:g} & {fmt_pm(a['f1'], a['f1_u'])} & {fmt_pm(a['f2'], a['f2_u'])} & {fmt_pm(a['f_avg'], a['f_avg_u'])} & {fmt_pm(a['f_beat'], a['f_beat_u'])} & {fmt_pm(a['y'], a['y_u'])} & {fmt_pm(a['ln_d'], a['ln_d_u'])} & {fmt_pm(a['ln_y'], a['ln_y_u'])} \\\\")
    L += [r"\bottomrule", r"\end{tabular}"]; open(os.path.join(args.out, "tables", "processed.tex"), "w").write("\n".join(L))
    if fit:
        L = [r"\begin{tabular}{@{}lll@{}}", r"\toprule", r"quantity & value & method \\", r"\midrule",
             rf"gradient $-n$ & ${-fit['n_best']:.2f} \pm {fit['dn_maxmin']:.2f}$ & best fit; max/min gradient lines ({-fit['n_max']:.2f}, {-fit['n_min']:.2f}) \\",
             rf"gradient $-n$ & ${-fit['n_best']:.3f} \pm {fit['se_n']:.3f}$ & LINEST standard error \\",
             rf"$R^2$ & {fit['R2']:.4f} & \\",
             rf"$\ln C$ & ${fit['lnC']:.2f} \pm {fit['se_lnC']:.2f}$ & intercept \\",
             rf"$C$ / m$^5$ s$^{{-2}}$ & ${fit['C']:.2e}$ & from intercept (free $n$) \\",
             rf"$C$ / m$^5$ s$^{{-2}}$ & {fmt_pm(fit['C5_mean'], fit['C5_hr'])} & mean of $(f_2^2-f_1^2)d^5$, fixed $n=5$ \\",
             rf"$\bar f_1$ / Hz & {fmt_pm(fit['f1_mean'], fit['f1_hr'])} & mean and half-range over all $d$ \\",
             rf"slope of $f_1$ vs $d$ / mHz mm$^{{-1}}$ & ${fit['f1_slope']*1e3:.2f} \pm {fit['f1_slope_se']*1e3:.2f}$ & linear regression \\",
             r"\bottomrule", r"\end{tabular}"]
        open(os.path.join(args.out, "tables", "fit.tex"), "w").write("\n".join(L))

    TEX = False
    # ---------- Excel workbook
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import ScatterChart, Reference, Series
    from openpyxl.chart.error_bar import ErrorBars
    from openpyxl.chart.data_source import NumDataSource, NumRef
    from openpyxl.chart.marker import Marker
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook(); bold = Font(bold=True); head_fill = PatternFill("solid", fgColor="DDEBF7")
    def put_table(ws, r0, header, rows, widths=None):
        for j, h in enumerate(header, 1):
            c = ws.cell(row=r0, column=j, value=h); c.font = bold; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and not np.isfinite(v): v = None
                ws.cell(row=r0 + i, column=j, value=(float(v) if isinstance(v, (np.floating, np.integer)) else v))
        if widths:
            for j, wdt in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = wdt
        return r0 + len(rows)
    def rng(ws, col, r0, r1): return f"'{ws.title}'!${get_column_letter(col)}${r0}:${get_column_letter(col)}${r1}"
    def scatter(ws, anchor, title, xlab, ylab, specs, w=18, h=10, xmin=None, xmax=None, ymin=None, ymax=None, xlog=False, ylog=False):
        ch = ScatterChart(); ch.title = title; ch.style = 13; ch.x_axis.title = xlab; ch.y_axis.title = ylab; ch.width = w; ch.height = h
        ch.x_axis.delete = False; ch.y_axis.delete = False
        if xmin is not None: ch.x_axis.scaling.min = xmin
        if xmax is not None: ch.x_axis.scaling.max = xmax
        if ymin is not None: ch.y_axis.scaling.min = ymin
        if ymax is not None: ch.y_axis.scaling.max = ymax
        if xlog: ch.x_axis.scaling.logBase = 10
        if ylog: ch.y_axis.scaling.logBase = 10
        for sp in specs:
            s = Series(Reference(sp["ws"], min_col=sp["y"], min_row=sp["r0"], max_row=sp["r1"]), Reference(sp["ws"], min_col=sp["x"], min_row=sp["r0"], max_row=sp["r1"]), title=sp["name"])
            if sp.get("line"):
                s.marker = Marker(symbol="none"); s.graphicalProperties.line.width = 15000; s.smooth = sp.get("smooth", True)
                if sp.get("dash"): s.graphicalProperties.line.dashStyle = sp["dash"]
                if sp.get("color"): s.graphicalProperties.line.solidFill = sp["color"]
            else:
                s.marker = Marker(symbol=sp.get("sym", "circle"), size=6); s.graphicalProperties.line.noFill = True
                if sp.get("color"): s.marker.graphicalProperties.solidFill = sp["color"]; s.marker.graphicalProperties.line.solidFill = sp["color"]
            if sp.get("ey"):
                f = rng(sp["ws"], sp["ey"], sp["r0"], sp["r1"])
                s.errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust", noEndCap=False, plus=NumDataSource(numRef=NumRef(f=f)), minus=NumDataSource(numRef=NumRef(f=f)))
            ch.series.append(s)
        ws.add_chart(ch, anchor)

    # -- sheet: Notes
    ws = wb.active; ws.title = "Notes"
    notes = ["Coupled magnetic-mechanical oscillator - data processing (generated by analyse_folder.py)", "",
             f"Input folder: {os.path.abspath(args.folder)}   ({len(trials)} trials, {len(agg)} separations{', ' + str(len(inf_rows)) + ' Infinity recordings' if inf_rows else ''})",
             f"Units: x converted from {args.unit} to mm; frequencies in Hz; d in mm (uncertainty +-{args.dd} mm, ruler/vernier reading + magnet mounting)", "",
             "METHOD PER TRIAL", "1. Release detected automatically: the flat 'hold' segment is removed; the analysis starts when both masses are moving (column t_start).",
             "2. Means removed (equilibrium positions); normal coordinates eta = xA + xB (in-phase) and xi = xB - xA (anti-phase) formed.",
             f"3. |FFT| of eta and xi ({args.window} window, zero-padded to 2^18 points, parabolic peak interpolation): f1 = peak of eta, f2 = peak of xi.",
             "   Each normal coordinate contains only ONE frequency, so the two peaks never overlap (columns eta_leak / xi_leak show the residual of the other mode; ~0 = identical springs).",
             "4. Cross-checks: two-peak FFT of xA alone and xB alone (Hann); beat period from the envelope maxima; peak-height ratio vs (f1/f2)^2; decay time tau.",
             "", "UNCERTAINTIES (international-essay / IB style)",
             "- FFT resolution: one frequency bin = 1/T; reading uncertainty of a peak taken as half a bin = 1/(2T) (column halfbin).",
             "- Random: half-range (max - min)/2 of the 5 trials (sample too small for a reliable standard deviation; the SD is listed for comparison).",
             f"- Adopted uncertainty of f1, f2 at each d: {dict(max='the larger of half-range and half-bin', halfrange='half-range', halfbin='half-bin', sd='standard deviation')[args.unc]}.",
             "- Propagation (absolute uncertainties added): d(f_avg) = (df1 + df2)/2;  d(f_beat) = df1 + df2;  d(f2^2 - f1^2) = 2 f2 df2 + 2 f1 df1;  d(ln y) = dy / y;  d(ln d) = dd / d.",
             "- Gradient of the ln-ln graph: best fit (least squares, = Excel LINEST) with the max/min gradient lines through the extreme error bars of the first and last points; dn = (n_max - n_min)/2.",
             "", "SHEETS", "Trials: every file.  Raw f1 / Raw f2: trial values, mean +- uncertainty (Table 4 style).  Processed: derived quantities with uncertainties (Table 5 style).",
             "Fits: gradient, C, f1 constancy, theory-vs-experiment comparison (Table 7 style).  Charts: native Excel charts (y error bars) + the matplotlib graphs (x and y error bars).",
             "Windows: the whole analysis repeated with rectangular / Hann / Hamming / Blackman windows and with both extraction methods (normal coordinates; conventional two-peak FFT of mass A) -> Graphs 6 and 7.",
             "Sample calc: live Excel formulas reproducing the processing for the first separation.  Infinity: uncoupled recordings (fa, fb) if present.",
             "", "WARNINGS PRINTED DURING THE RUN"] + [l for l in log.getvalue().splitlines() if "WARNING" in l]
    for i, line in enumerate(notes, 1): ws.cell(row=i, column=1, value=line).font = bold if (i == 1 or line.isupper()) else Font()
    ws.column_dimensions["A"].width = 150

    # -- sheet: Trials
    ws = wb.create_sheet("Trials")
    tcols = ["d_mm", "trial", "file", "fs", "T", "N", "t_start", "t_start_A", "t_start_B", "B_preswing_mm", "bin_Hz", "halfbin_Hz", "method", "f1", "f2", "f1_winspread", "f2_winspread",
             "f1_A", "f2_A", "f1_B", "f2_B", "eta_leak", "xi_leak", "rA", "rB", "ratio_pred", "T_beat_fft", "T_beat_env", "dT_beat_env", "n_beats", "env_minmax", "tau_s", "d_track_mm", "A_mm"]
    put_table(ws, 1, tcols, [[r.get(c, np.nan) for c in tcols] for r in trials], widths=[8, 6, 16] + [10] * 30); ws.freeze_panes = "D2"

    # -- sheets: Raw f1, Raw f2
    for key, label in (("f1", "f1"), ("f2", "f2")):
        ws = wb.create_sheet(f"Raw {label}")
        header = ["d / mm", "dd / mm"] + [f"{label} trial {i+1} / Hz" for i in range(ntr)] + ["mean / Hz", "half-range / Hz", "SD / Hz", "T mean / s", "half-bin 1/(2T) / Hz", "window spread / Hz", f"adopted d{label} / Hz"]
        rows = [[a["d_mm"], args.dd] + a[key + "_trials"] + [np.nan] * (ntr - len(a[key + "_trials"])) + [a[key], a[key + "_hr"], a[key + "_sd"], a["T_mean"], a["halfbin"], a[key + "_ws"], a[key + "_u"]] for a in agg]
        put_table(ws, 1, header, rows, widths=[9, 8] + [13] * ntr + [11, 13, 10, 10, 16, 15, 14])
        for r in range(2, 2 + len(rows)):
            for c in range(3, 3 + ntr + 7): ws.cell(row=r, column=c).number_format = "0.0000"

    # -- sheet: Processed
    ws = wb.create_sheet("Processed")
    header = ["d / mm", "dd / mm", "f1 / Hz", "df1 / Hz", "f2 / Hz", "df2 / Hz", "f_avg / Hz", "df_avg / Hz", "f_beat / Hz", "df_beat / Hz", "f2^2-f1^2 / Hz^2", "d(f2^2-f1^2) / Hz^2",
              "ln(d/m)", "d ln d", "ln(f2^2-f1^2)", "d ln(...)", "C_i=(f2^2-f1^2)d^5 / m^5 s^-2", "dC_i", "2gamma/k=(f2/f1)^2-1", "(f1/f2)^2", "rA", "rB", "tau / s", "d_track / mm"]
    rows = [[a["d_mm"], args.dd, a["f1"], a["f1_u"], a["f2"], a["f2_u"], a["f_avg"], a["f_avg_u"], a["f_beat"], a["f_beat_u"], a["y"], a["y_u"], a["ln_d"], a["ln_d_u"], a["ln_y"], a["ln_y_u"],
             a["C_i"], a["C_i_u"], a["coupling"], a["ratio_pred"], a["rA"], a["rB"], a["tau"], a["d_track"]] for a in agg]
    put_table(ws, 1, header, rows, widths=[9, 8] + [11] * 22); nP = len(rows); P = ws
    for r in range(2, 2 + nP):
        for c in range(3, 25): ws.cell(row=r, column=c).number_format = "0.0000" if c < 17 else "0.00E+00" if c < 19 else "0.000"

    # -- sheet: Curves (for the chart lines) and Fits
    if fit:
        wc = wb.create_sheet("Curves"); dsm_mm = np.linspace(0.9 * dmm.min(), 1.08 * dmm.max(), 60); dsm = dsm_mm / 1e3
        lnx = np.linspace(fit["X"].min() - 0.05, fit["X"].max() + 0.05, 30)
        rows = [[dsm_mm[i], math.sqrt(f1m ** 2 + fit["C5_ls"] * dsm[i] ** -5), f1m, math.sqrt(f1m ** 2 + fit["C5_ls"] * dsm[i] ** -5) - f1m, fit["C5_ls"] * dsm[i] ** -5,
                 lnx[i] if i < 30 else None, fit["lnC"] - fit["n_best"] * lnx[i] if i < 30 else None, fit["c_steep"] - fit["n_max"] * lnx[i] if i < 30 else None, fit["c_shallow"] - fit["n_min"] * lnx[i] if i < 30 else None,
                 fit["C5_ls"]] for i in range(60)]
        put_table(wc, 1, ["d / mm", "f2 model (n=5) / Hz", "f1 mean / Hz", "f_beat model / Hz", "f2^2-f1^2 model / Hz^2", "ln d", "ln y best fit", "ln y max gradient", "ln y min gradient", "C (n=5)"], rows, widths=[10] * 10)
        wf = wb.create_sheet("Fits")
        rows = [["ln-ln graph: gradient -n (best fit)", -fit["n_best"], "least squares (Excel LINEST)"], ["standard error of gradient", fit["se_n"], "LINEST"], ["max gradient", -fit["n_max"], "line through extreme error-bar corners"],
                ["min gradient", -fit["n_min"], ""], ["n = ", f"{fit['n_best']:.2f} ± {fit['dn_maxmin']:.2f}", "(n_max - n_min)/2"], ["R^2", fit["R2"], ""],
                ["intercept ln C", fit["lnC"], f"± {fit['se_lnC']:.3f} (LINEST)"], ["C from intercept / m^5 s^-2", fit["C"], f"± {fit['C_u_maxmin']:.2e} (max/min lines)"],
                ["", None, ""], ["fixed n = 5: C (least squares) / m^5 s^-2", fit["C5_ls"], ""], ["fixed n = 5: mean of C_i", fit["C5_mean"], f"half-range {fit['C5_hr']:.2e}, SD {fit['C5_sd']:.2e}"],
                ["chi^2 / dof for n = 5", f"{fit['chi2_n5']:.1f} / {fit['dof_n5']}", "with the adopted uncertainties"], ["", None, ""],
                ["f1: mean over all d / Hz", fit["f1_mean"], f"half-range {fit['f1_hr']:.4f}, SD {fit['f1_sd']:.4f}"], ["f1: slope against d / Hz per mm", fit["f1_slope"], f"± {fit['f1_slope_se']:.2e} (consistent with 0 => f1 independent of d)"],
                ["f1: largest deviation from mean / Hz", fit["f1_max_dev"], f"mean adopted uncertainty {fit['f1_mean_unc']:.4f} Hz"]]
        r_end = put_table(wf, 1, ["quantity", "value", "note"], rows, widths=[42, 22, 60])
        for r in range(2, r_end + 1):
            v = wf.cell(row=r, column=2).value
            if isinstance(v, float) and 0 < abs(v) < 1e-3: wf.cell(row=r, column=2).number_format = "0.000E+00"
        # theory vs experiment (Table 7 style)
        r0 = r_end + 2; wf.cell(row=r0, column=1, value="Comparison with the dipole model f2 = sqrt(f1_mean^2 + C d^-5), C from the fixed-n fit").font = bold
        rows = []
        for a in agg:
            f2th = math.sqrt(f1m ** 2 + fit["C5_ls"] * (a["d_mm"] / 1e3) ** -5); rows.append([a["d_mm"], a["f2"], a["f2_u"], f2th, 100 * (a["f2"] - f2th) / f2th, "yes" if abs(a["f2"] - f2th) <= a["f2_u"] else "no"])
        put_table(wf, r0 + 1, ["d / mm", "f2 exp / Hz", "df2 / Hz", "f2 theory / Hz", "% difference", "within error bar?"], rows)

    # -- sheet: Sample calc (live formulas for the first separation)
    ws = wb.create_sheet("Sample calc"); a0 = agg[0]; n0 = len(a0["f1_trials"]); L = get_column_letter
    ws["A1"] = f"Sample calculation for d = {a0['d_mm']} mm (live Excel formulas; same steps as the Python code)"; ws["A1"].font = bold
    ws["A3"] = "trial"; ws["A4"] = "f1 / Hz"; ws["A5"] = "f2 / Hz"; ws["A6"] = "T / s"
    for i in range(n0):
        ws.cell(row=3, column=2 + i, value=i + 1); ws.cell(row=4, column=2 + i, value=float(a0["f1_trials"][i])); ws.cell(row=5, column=2 + i, value=float(a0["f2_trials"][i]))
        ws.cell(row=6, column=2 + i, value=float([r["T"] for r in trials if r["d_mm"] == a0["d_mm"]][i]))
    c1, cn = L(2), L(1 + n0); R = 8
    lines = [("mean f1", f"=AVERAGE({c1}4:{cn}4)"), ("half-range f1", f"=(MAX({c1}4:{cn}4)-MIN({c1}4:{cn}4))/2"), ("mean T", f"=AVERAGE({c1}6:{cn}6)"), ("half-bin = 1/(2T)", "=0.5/B10"),
             ("adopted df1 = MAX(half-range, half-bin)", "=MAX(B9,B11)"), ("mean f2", f"=AVERAGE({c1}5:{cn}5)"), ("half-range f2", f"=(MAX({c1}5:{cn}5)-MIN({c1}5:{cn}5))/2"), ("adopted df2", "=MAX(B14,B11)"),
             ("f_avg = (f1+f2)/2", "=(B8+B13)/2"), ("df_avg = (df1+df2)/2", "=(B12+B15)/2"), ("f_beat = f2-f1", "=B13-B8"), ("df_beat = df1+df2", "=B12+B15"),
             ("y = f2^2-f1^2", "=B13^2-B8^2"), ("dy = 2 f2 df2 + 2 f1 df1", "=2*B13*B15+2*B8*B12"), ("ln(d/m)", f"=LN({a0['d_mm']}/1000)"), ("d ln d = dd/d", f"={args.dd}/{a0['d_mm']}"),
             ("ln y", "=LN(B20)"), ("d ln y = dy/y", "=B21/B20"), ("C_i = y d^5 / m^5 s^-2", f"=B20*({a0['d_mm']}/1000)^5"), ("(f1/f2)^2 (predicted peak-height ratio)", "=(B8/B13)^2")]
    for i, (lab, frm) in enumerate(lines):
        ws.cell(row=R + i, column=1, value=lab); ws.cell(row=R + i, column=2, value=frm)
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 16
    if winfit:
        ww = wb.create_sheet("Windows"); WN = ("rect", "hann", "hamming", "blackman"); r0 = 1
        for m, mlab in (("", "Normal-coordinate method: f1 = peak of FFT(xA+xB), f2 = peak of FFT(xB-xA)"), ("A", "Conventional method: the two peaks of FFT(xA) alone")):
            ww.cell(row=r0, column=1, value=mlab).font = bold
            header = ["d / mm"] + [f"{c} ({wn})" for wn in WN for c in ("f1 / Hz", "df1", "f2 / Hz", "df2", "f2^2-f1^2 / Hz^2", "d(f2^2-f1^2)")]
            rows = [[a["d_mm"]] + [a[k] for wn in WN for k in (f"f1{m}_{wn}", f"f1{m}_{wn}_u", f"f2{m}_{wn}", f"f2{m}_{wn}_u", f"y{m}_{wn}", f"y{m}_{wn}_u")] for a in agg]
            r_end = put_table(ww, r0 + 1, header, rows, widths=[9] + [12] * 24); nW = len(rows); first_row = r0 + 2
            put_table(ww, r_end + 2, ["window", "gradient -n", "± (max/min)", "s.e. (LINEST)", "R^2", "C / m^5 s^-2", "mean f1 / Hz", "half-range f1 / Hz"],
                      [[wn, -winfit[m + wn]["n"], winfit[m + wn]["dn"], winfit[m + wn]["se"], winfit[m + wn]["R2"], winfit[m + wn]["C"], winfit[m + wn]["f1"], winfit[m + wn]["f1_hr"]] for wn in WN])
            for r in range(r_end + 3, r_end + 7): ww.cell(row=r, column=6).number_format = "0.000E+00"
            scatter(ww, get_column_letter(27) + str(r0), f"f2 against d, each window ({'normal coordinates' if m == '' else 'FFT of xA alone'})", "d / mm", "f2 / Hz",
                    [dict(ws=ww, x=1, y=4 + 6 * i, ey=5 + 6 * i, r0=first_row, r1=first_row + nW - 1, name=WN[i], sym=("circle", "square", "triangle", "diamond")[i], color=("000000", "C00000", "1F4E79", "548235")[i]) for i in range(4)], w=18, h=9)
            r0 = r_end + 10
    if inf_rows:
        ws = wb.create_sheet("Infinity"); put_table(ws, 1, ["file", "T / s", "fs / Hz", "fA / Hz", "fB / Hz", "half-bin / Hz", "tau / s"], [[r["file"], r["T"], r["fs"], r["f1_A"], r["f1_B"], r["halfbin_Hz"], r["tau_s"]] for r in inf_rows], widths=[18, 8, 8, 10, 10, 10, 8])

    # -- sheet: Charts
    wsC = wb.create_sheet("Charts")
    if fit:
        wc = wb["Curves"]
        scatter(wsC, "A1", "f1 and f2 against d", "d / mm", "frequency / Hz",
                [dict(ws=P, x=1, y=5, ey=6, r0=2, r1=1 + nP, name="f2 measured", color="C00000"), dict(ws=P, x=1, y=3, ey=4, r0=2, r1=1 + nP, name="f1 measured", color="B35900", sym="diamond"),
                 dict(ws=wc, x=1, y=2, r0=2, r1=61, name="f2 dipole model (C fitted, n = 5)", line=True, color="C00000"), dict(ws=wc, x=1, y=3, r0=2, r1=61, name="mean f1", line=True, dash="dash", color="B35900")], ymin=math.floor(f1m * 2) / 2)
        scatter(wsC, "L1", "ln(f2^2 - f1^2) against ln d", "ln(d / m)", "ln[(f2^2 - f1^2) / Hz^2]",
                [dict(ws=P, x=13, y=15, ey=16, r0=2, r1=1 + nP, name="data", color="000000"), dict(ws=wc, x=6, y=7, r0=2, r1=31, name=f"best fit, gradient {-fit['n_best']:.2f}", line=True, smooth=False, color="C00000"),
                 dict(ws=wc, x=6, y=8, r0=2, r1=31, name=f"max gradient {-fit['n_max']:.2f}", line=True, smooth=False, dash="dash", color="7F7F7F"), dict(ws=wc, x=6, y=9, r0=2, r1=31, name=f"min gradient {-fit['n_min']:.2f}", line=True, smooth=False, dash="sysDot", color="7F7F7F")])
        scatter(wsC, "A22", "f_beat = f2 - f1 against d", "d / mm", "f_beat / Hz",
                [dict(ws=P, x=1, y=9, ey=10, r0=2, r1=1 + nP, name="measured", color="C00000"), dict(ws=wc, x=1, y=4, r0=2, r1=61, name="dipole model", line=True, color="C00000")])
        scatter(wsC, "L22", "(f2^2 - f1^2) d^5 against d  (flat = d^-5 law)", "d / mm", "C_i / m^5 s^-2",
                [dict(ws=P, x=1, y=17, ey=18, r0=2, r1=1 + nP, name="C_i", color="000000"), dict(ws=wc, x=1, y=10, r0=2, r1=61, name="C (fixed n = 5)", line=True, color="C00000")], ymin=0)
        scatter(wsC, "A43", "f_avg against d", "d / mm", "f_avg / Hz", [dict(ws=P, x=1, y=7, ey=8, r0=2, r1=1 + nP, name="measured", color="1F4E79")])
        scatter(wsC, "L43", "FFT peak-height ratio against d", "d / mm", "|X(f2)| / |X(f1)|",
                [dict(ws=P, x=1, y=21, r0=2, r1=1 + nP, name="mass A", color="B35900"), dict(ws=P, x=1, y=22, r0=2, r1=1 + nP, name="mass B", color="1F4E79", sym="square"), dict(ws=P, x=1, y=20, r0=2, r1=1 + nP, name="(f1/f2)^2 predicted", line=True, smooth=False, color="000000")], ymin=0)
        row = 64
        for g in ("graph1_frequencies", "graph2_loglog", "graph3_beat_avg", "graph4_exponent_test", "graph5_peak_ratio", "graph6_windows", "graph7_window_shift"):
            p = os.path.join(args.out, g + ".png")
            if os.path.exists(p):
                img = XLImage(p); img.width, img.height = img.width * 0.42, img.height * 0.42; wsC.add_image(img, f"A{row}"); row += int(img.height / 20) + 2
    wb.save(os.path.join(args.out, "results.xlsx"))
    open(os.path.join(args.out, "summary.txt"), "w").write(log.getvalue())
    say(f"\nwritten: {args.out}/results.xlsx, trials.csv, graph1-7 (.png/.pdf), tables/*.tex, figs/trial_*.png, summary.txt")

if __name__ == "__main__":
    main()
